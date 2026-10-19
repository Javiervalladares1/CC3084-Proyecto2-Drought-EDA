"""Exportadores Word y PDF de una única fuente Markdown."""
from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, PageBreak, Paragraph, Spacer, Table, TableStyle, Image

from crear_revision_word import configure, rich

ROOT = Path(__file__).resolve().parents[1]


def fractions(headers):
    count = len(headers)
    return {2: [.36, .64], 3: [.27, .36, .37], 4: [.25, .25, .25, .25],
            5: [.28, .18, .18, .18, .18], 6: [.25, .15, .15, .15, .15, .15]}[count]


def blocks(path):
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    i, code = 0, False
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if line.startswith("```"):
            code = not code
            continue
        if not line:
            continue
        if line == "<!-- pagebreak -->":
            yield "break", None
        elif line.startswith("|"):
            table_lines = [line]
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [[value.strip() for value in row.strip("|").split("|")] for row in table_lines]
            yield "table", [rows[0], *rows[2:]]
        elif line.startswith("!["):
            match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            yield "image", (match[1], (Path(path).parent / match[2]).resolve())
        elif line.startswith("#"):
            match = re.match(r"^(#{1,3})\s+(.*)$", line)
            yield "heading", (len(match[1]), match[2])
        elif code:
            yield "code", line
        else:
            yield "paragraph", line


def word_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    ratios = fractions(rows[0])
    for column, ratio in zip(table.columns, ratios):
        column.width = Cm(16.7 * ratio)
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        for key, value in {"val": "single", "sz": "4", "color": "000000"}.items():
            el.set(qn("w:" + key), value)
        borders.append(el)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", 85), ("bottom", 85), ("left", 100), ("right", 100)):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    props.append(margins)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, row_values in enumerate(rows):
        row = table.rows[i]
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, value in enumerate(row_values):
            cell = row.cells[j]
            cell.width = Cm(16.7 * ratios[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = "Table Text"
            if i > 0 and re.fullmatch(r"[-+\d., %]+", value):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rich(p, value)
            for run in p.runs:
                run.font.size = Pt(8.9)
                if i == 0:
                    run.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(5)


def export_word(source, output):
    doc = Document()
    configure(doc)
    doc.sections[0].header.paragraphs[0].runs[0].text = "CC3084 · PROYECTO 2 · RESULTADOS"
    doc.core_properties.title = "Pronóstico del almacenamiento total de agua Resultados del Proyecto 2"
    doc.core_properties.author = "Javier Valladares, Ian Cumes y Javier Cifuentes"
    doc.core_properties.subject = "Random Forest, XGBoost y redes LSTM con evaluación temporal"
    pending_break, page, references = False, 1, False
    for kind, value in blocks(source):
        if kind == "break":
            pending_break, page = True, page + 1
            continue
        if kind == "heading":
            level, text = value
            if text.startswith(("Referencias", "Fuentes de documentación")):
                references = True
            elif text.startswith("Anexo"):
                references = False
            style = "Title" if page == 1 and level == 1 else "Subtitle" if page == 1 and level == 2 else f"Heading {level}"
            p = doc.add_paragraph(style=style)
            rich(p, text)
        elif kind == "table":
            word_table(doc, value)
            continue
        elif kind == "image":
            alt, path = value
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Cm(16.5))
            # La leyenda es texto editable y viaja junto a su figura.
            caption = doc.add_paragraph(style="Note")
            rich(caption, alt)
        else:
            is_note = value.startswith("*Nota")
            is_caption = value.startswith("**Tabla")
            style = "Note" if is_note else "Reference" if references else "Normal"
            p = doc.add_paragraph(style=style)
            if style == "Normal":
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if is_caption:
                p.paragraph_format.keep_with_next = True
            if kind == "code":
                p.paragraph_format.line_spacing = Pt(13)
                run = p.add_run(value)
                run.font.name = "Courier New"
                run.font.size = Pt(9)
            else:
                rich(p, value)
        if pending_break:
            p.paragraph_format.page_break_before = True
            pending_break = False
    for part in doc.part.package.parts:
        element = getattr(part, "element", None)
        if element is None and str(part.partname) == "/word/stylesWithEffects.xml":
            element = etree.fromstring(part.blob)
            for color in element.iter(qn("w:color")):
                color.attrib.clear()
                color.set(qn("w:val"), "000000")
            part._blob = etree.tostring(element, xml_declaration=True, encoding="UTF-8", standalone=True)
        elif element is not None:
            for color in element.xpath(".//w:color"):
                color.attrib.clear()
                color.set(qn("w:val"), "000000")
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def pdf_rich(text):
    text = html.escape(text)
    text = re.sub(r"\[([^\]]+)\]\((https?://[^\s)]+)\)", lambda m: f'<link href="{m[2]}" color="#000000">{m[1]}</link>', text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    return re.sub(r"\*([^*]+)\*", r"<i>\1</i>", text)


def export_pdf(source, output):
    directory = Path("/System/Library/Fonts/Supplemental")
    if directory.exists():
        for name, filename in [("Results", "Arial.ttf"), ("Results-Bold", "Arial Bold.ttf"),
                               ("Results-Italic", "Arial Italic.ttf"), ("Results-BoldItalic", "Arial Bold Italic.ttf")]:
            pdfmetrics.registerFont(TTFont(name, str(directory / filename)))
        pdfmetrics.registerFontFamily("Results", normal="Results", bold="Results-Bold", italic="Results-Italic", boldItalic="Results-BoldItalic")
        font, bold = "Results", "Results-Bold"
    else:
        font, bold = "Helvetica", "Helvetica-Bold"
    black = colors.black
    styles = {
        "body": ParagraphStyle("Body", fontName=font, fontSize=10.8, leading=15, textColor=black, alignment=TA_JUSTIFY, spaceAfter=9),
        "title": ParagraphStyle("Title", fontName=bold, fontSize=25, leading=31, textColor=black, spaceAfter=15),
        "subtitle": ParagraphStyle("Subtitle", fontName=font, fontSize=16, leading=21, textColor=black, spaceAfter=22),
        "h1": ParagraphStyle("H1", fontName=bold, fontSize=17, leading=21, textColor=black, spaceAfter=17, keepWithNext=True),
        "h2": ParagraphStyle("H2", fontName=bold, fontSize=12.5, leading=17, textColor=black, spaceAfter=9, spaceBefore=5, keepWithNext=True),
        "h3": ParagraphStyle("H3", fontName=bold, fontSize=11.2, leading=15, textColor=black, spaceAfter=6, spaceBefore=5, keepWithNext=True),
        "note": ParagraphStyle("Note", fontName=font, fontSize=9.2, leading=12.7, textColor=black, spaceAfter=10),
        "ref": ParagraphStyle("Reference", fontName=font, fontSize=10.6, leading=20.8, textColor=black, alignment=TA_LEFT,
                              leftIndent=1.27 * cm, firstLineIndent=-1.27 * cm, spaceAfter=10),
        "cell": ParagraphStyle("Cell", fontName=font, fontSize=8.9, leading=12, textColor=black),
        "cellhead": ParagraphStyle("CellHead", fontName=bold, fontSize=8.9, leading=12, textColor=black),
        "code": ParagraphStyle("Code", fontName="Courier", fontSize=9, leading=13, textColor=black, spaceAfter=5),
    }
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(output), pagesize=A4, leftMargin=2.15*cm, rightMargin=2.15*cm,
                            topMargin=1.95*cm, bottomMargin=1.9*cm)
    story, page, references = [], 1, False
    for kind, value in blocks(source):
        if kind == "break":
            story.append(PageBreak())
            page += 1
        elif kind == "heading":
            level, text = value
            if text.startswith(("Referencias", "Fuentes de documentación")):
                references = True
            elif text.startswith("Anexo"):
                references = False
            style = "title" if page == 1 and level == 1 else "subtitle" if page == 1 and level == 2 else f"h{level}"
            story.append(Paragraph(pdf_rich(text), styles[style]))
        elif kind == "table":
            cells = [[Paragraph(pdf_rich(text), styles["cellhead" if i == 0 else "cell"]) for text in row] for i, row in enumerate(value)]
            table = Table(cells, colWidths=[doc.width*f for f in fractions(value[0])], repeatRows=1)
            table.setStyle(TableStyle([("VALIGN", (0,0),(-1,-1),"MIDDLE"), ("GRID",(0,0),(-1,-1),.4,black),
                                       ("LEFTPADDING",(0,0),(-1,-1),5), ("RIGHTPADDING",(0,0),(-1,-1),5),
                                       ("TOPPADDING",(0,0),(-1,-1),6), ("BOTTOMPADDING",(0,0),(-1,-1),6)]))
            story.extend([table, Spacer(1, 9)])
        elif kind == "image":
            alt, path = value
            with PILImage.open(path) as image:
                width, height = image.size
            display_width = 16.5 * cm
            story.extend([Image(str(path), width=display_width, height=display_width * height/width), Spacer(1, 5), Paragraph(pdf_rich(alt), styles["note"])])
        else:
            style = "code" if kind == "code" else "note" if value.startswith("*Nota") else "ref" if references else "body"
            story.append(Paragraph(pdf_rich(value), styles[style]))
    def decorate(canvas, current):
        canvas.saveState()
        canvas.setFillColor(black)
        canvas.setFont(font, 8)
        canvas.drawString(doc.leftMargin, A4[1] - .95*cm, "CC3084 · PROYECTO 2 · RESULTADOS")
        canvas.drawRightString(A4[0]-doc.rightMargin, .85*cm, f"Página {current.page}")
        canvas.setTitle("Pronóstico del almacenamiento total de agua Resultados del Proyecto 2")
        canvas.setAuthor("Javier Valladares, Ian Cumes y Javier Cifuentes")
        canvas.restoreState()
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
