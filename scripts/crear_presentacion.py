"""
Genera la presentación de resultados del Proyecto 2.

La presentación está diseñada para durar 15 minutos repartidos entre cuatro
expositores. Cada diapositiva lleva notas del orador con el tiempo asignado y
con lo que conviene decir, de modo que el guion viaje dentro del propio archivo.

El contenido proviene del análisis del notebook y las figuras de figures/.

Uso (desde la raíz del proyecto):
    python3 scripts/crear_presentacion.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

RAIZ = Path(__file__).resolve().parent.parent
FIGURAS = RAIZ / "figures"
DESTINO = RAIZ / "docs" / "Proyecto2_Presentacion.pptx"

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0F, 0x6E, 0x63)
GRIS = RGBColor(0x44, 0x4A, 0x52)

ANCHO, ALTO = Inches(13.333), Inches(7.5)

# Registro de tiempos para el resumen que imprime el script.
guion = []


def nueva(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # en blanco


def titulo(slide, texto):
    caja = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), ANCHO - Inches(1.2), Inches(0.8))
    p = caja.text_frame.paragraphs[0]
    p.text = texto
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = AZUL


def vinetas(slide, items, izquierda=0.7, arriba=1.3, ancho=None, tam=16):
    ancho = ancho or (ANCHO - Inches(1.4))
    caja = slide.shapes.add_textbox(Inches(izquierda), Inches(arriba), ancho, Inches(4.4))
    tf = caja.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        sangrado = item.startswith("    ")
        p.text = item if sangrado else f"•  {item}"
        p.font.size = Pt(tam - 2 if sangrado else tam)
        p.font.color.rgb = GRIS
        p.space_after = Pt(8)


def mensaje(slide, texto):
    caja = slide.shapes.add_textbox(Inches(0.7), ALTO - Inches(0.78), ANCHO - Inches(1.4), Inches(0.55))
    p = caja.text_frame.paragraphs[0]
    p.text = texto
    p.font.size = Pt(15)
    p.font.italic = True
    p.font.bold = True
    p.font.color.rgb = TEAL


def imagen(slide, nombre, izquierda, arriba, ancho):
    ruta = FIGURAS / nombre
    if ruta.exists():
        slide.shapes.add_picture(str(ruta), Inches(izquierda), Inches(arriba), width=Inches(ancho))


def notas(slide, expositor, segundos, texto):
    """Escribe el guion en las notas del orador y registra el tiempo."""
    guion.append((len(guion) + 1, expositor, segundos))
    minutos = f"{segundos // 60}:{segundos % 60:02d}"
    slide.notes_slide.notes_text_frame.text = (
        f"[{expositor} · {minutos}]\n\n{texto.strip()}"
    )


prs = Presentation()
prs.slide_width, prs.slide_height = ANCHO, ALTO

E1, E2, E3, E4 = "Integrante 1", "Integrante 2", "Integrante 3", "Integrante 4"

# =============================================================================
# BLOQUE 1 — El problema
# =============================================================================

# --- 1. Portada -------------------------------------------------------------
s = nueva(prs)
caja = s.shapes.add_textbox(Inches(0.9), Inches(2.1), ANCHO - Inches(1.8), Inches(3))
tf = caja.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "A Step Ahead of Drought"
p.font.size = Pt(46); p.font.bold = True; p.font.color.rgb = AZUL
for texto, tam, color in [
    ("Pronóstico global de almacenamiento de agua · Reto ITU", 22, TEAL),
    ("Proyecto 2 — Análisis Exploratorio", 20, GRIS),
    ("CC3084 Data Science · Universidad del Valle de Guatemala · Semestre II 2026", 15, GRIS),
    ("Integrantes: ______________________", 14, GRIS),
    ("https://github.com/Javiervalladares1/CC3084-Proyecto2-Drought-EDA", 12, TEAL),
]:
    q = tf.add_paragraph(); q.text = texto
    q.font.size = Pt(tam); q.font.color.rgb = color
notas(s, E1, 15, """
Presentarse y presentar el reto en una frase: es una competencia de la ITU, alojada
en Zindi, sobre pronosticar el agua almacenada del planeta. Decir que en 15 minutos
se recorre el problema, los datos, los patrones encontrados y hacia dónde apunta el
modelado. No leer la diapositiva.
""")

# --- 2. Qué es el TWS -------------------------------------------------------
s = nueva(prs)
titulo(s, "¿Qué es el almacenamiento total de agua?")
vinetas(s, [
    "El TWS (Total Water Storage) es toda el agua de un territorio: subterránea, "
    "humedad del suelo, superficial, nieve y hielo.",
    "Las misiones satelitales GRACE y GRACE-FO no lo ven directamente: miden "
    "variaciones del campo gravitatorio terrestre.",
    "Como el agua tiene masa, un cambio en el almacenamiento cambia la gravedad "
    "local. De esa señal se deriva el TWS.",
    "Es el mejor indicador integrado de sequía hidrológica a gran escala, porque "
    "incluye el agua subterránea.",
], ancho=Inches(5.6), tam=15)
imagen(s, "18_mapa_promedio.png", 6.3, 2.0, 6.6)
mensaje(s, "Medimos el agua del planeta pesándolo desde el espacio.")
notas(s, E1, 60, """
Es el concepto clave: si no se entiende qué es el TWS, nada del resto se entiende.

Insistir en dos cosas. Primero, que no es solo lluvia: es el balance acumulado, e
incluye acuíferos. Segundo, el mecanismo de medición, que suele sorprender: el
satélite pesa la Tierra, no fotografía el agua.

Señalar en el mapa el norte de África y la India en café: son las zonas de menor
almacenamiento promedio. Ese mapa lo produjimos nosotros con los datos del reto.
""")

# --- 3. La latencia ---------------------------------------------------------
s = nueva(prs)
titulo(s, "El problema: la información llega tarde")
vinetas(s, [
    "La sequía se instala despacio. Cuando ya se ve en los cultivos o en los pozos, "
    "el déficit lleva meses acumulándose.",
    "Los productos derivados de GRACE se publican con un retraso de 2 a 3 meses.",
    "Para una alerta temprana se necesita el estado actual del sistema; lo que hay "
    "disponible es el estado de hace un trimestre.",
    "El reto pide predecir el TWS del mes siguiente a partir de la información del "
    "mes actual. La métrica de evaluación es el RMSE.",
], tam=17)
mensaje(s, "La información existe, pero llega tarde para sostener una decisión.")
notas(s, E1, 50, """
Aquí está la justificación del proyecto entero. Vale la pena decirlo despacio.

La idea que debe quedar: el problema no es que falte el dato, es que llega tarde.
Un buen ejemplo para dar contexto: si hoy quisiéramos saber si una región está
entrando en sequía, el dato satelital más reciente sería el de hace un trimestre.

Cerrar mencionando que la métrica es RMSE, porque volverá a aparecer al final
cuando hablemos de la línea base.
""")

# --- 4. Problema científico y objetivos -------------------------------------
s = nueva(prs)
titulo(s, "Problema científico y objetivos")
vinetas(s, [
    "Pregunta: ¿qué patrones temporales, estacionales y espaciales presenta el TWS, "
    "y qué variables del mes actual se asocian más con el valor del mes siguiente?",
    "Objetivo general: caracterizar esos patrones para sustentar un modelo posterior "
    "de pronóstico a un mes.",
    "Seis objetivos específicos: describir la estructura, evaluar la calidad, "
    "analizar el comportamiento temporal, comparar regiones, medir asociaciones e "
    "identificar variables candidatas.",
    "Esta etapa es exploratoria: el objetivo no era construir el modelo, sino "
    "entender el problema antes de modelarlo.",
], tam=16)
mensaje(s, "Una pregunta descriptiva y asociativa: en ningún momento afirmamos causalidad.")
notas(s, E1, 50, """
Leer la pregunta científica completa y pausado; es lo que la rúbrica evalúa.

Recalcar el último punto: es un EDA, no un modelo. Decidimos deliberadamente no
entrenar nada hasta entender los datos, y eso resultó acertado porque encontramos
cosas que habrían arruinado un modelo hecho a ciegas.

Cerrar con la advertencia de causalidad. Se repite al final, pero conviene
plantarla desde ya.

Aquí termina el primer bloque: dar paso al siguiente expositor.
""")

# =============================================================================
# BLOQUE 2 — Los datos
# =============================================================================

# --- 5. El dataset ----------------------------------------------------------
s = nueva(prs)
titulo(s, "Descripción del dataset")
vinetas(s, [
    "2 154 021 observaciones de entrenamiento y 280 961 de prueba, con 13 variables.",
    "Rejilla global de 1°, con 15 715 celdas, de −55.5° a 83.5° de latitud.",
    "Periodo de entrenamiento: mayo 2002 a agosto 2015 (138 meses observados).",
    "Formato plano: una fila por combinación de mes y celda.",
    "Variables: TWS_t, SPEI a 1/3/6/12 meses, humedad del suelo y la variable "
    "objetivo (TWS del mes siguiente).",
    "Todas están estandarizadas y son adimensionales: no son centímetros de agua.",
], tam=16)
mensaje(s, "Una fila por mes y por celda del planeta.")
notas(s, E2, 45, """
No leer los números uno por uno: son muchos. Quedarse con dos ideas.

Primera: es un dataset grande pero simple de forma. Una fila = un mes en una celda
de un grado.

Segunda, y es la que suele pasarse por alto: las variables vienen estandarizadas.
No podemos decir "se perdieron tantos centímetros de agua"; solo podemos hablar de
desviaciones relativas. Lo verificamos: la objetivo tiene media 0.11 y desviación
0.91, que es la firma de una variable estandarizada.
""")

# --- 6. Calidad -------------------------------------------------------------
s = nueva(prs)
titulo(s, "Calidad de los datos: lo que el conteo de nulos no ve")
vinetas(s, [
    "Cero valores nulos en las 13 columnas. Cero duplicados. Cero identificadores repetidos.",
    "Y sin embargo faltaba el 14 % del periodo.",
    "El formato es plano: si un mes no se observó, no genera filas — en vez de "
    "generar NaN. El conteo de nulos no puede detectarlo.",
    "Faltan 22 de los 160 meses del calendario, concentrados en 2002–2003 y "
    "2011–2014. Los años 2004 a 2010 están completos.",
], arriba=1.25, ancho=Inches(5.9), tam=15)
imagen(s, "10_cobertura_temporal.png", 6.5, 2.4, 6.3)
mensaje(s, "isna().sum() daba cero y aun así faltaba el 14 % del periodo.")
notas(s, E2, 80, """
Este es el hallazgo del que más orgullosos estamos. Darle tiempo.

Contarlo como historia: corrimos el chequeo estándar de calidad, salió impecable,
y por poco damos el dataset por bueno. Al comparar contra el calendario completo
apareció que faltaban 22 meses enteros.

La explicación del porqué es lo importante: en formato plano, un mes sin observar
desaparece en lugar de aparecer como nulo. Ninguna herramienta de perfilado
automático lo detecta.

Señalar en la gráfica las barras de 2011 a 2014, que pierden entre 4 y 5 meses
cada una. Coincide con los últimos años de operación de GRACE, aunque esa
atribución viene de literatura externa, no de nuestros datos.

Cerrar con la consecuencia práctica: al construir rezagos para el modelo, un
shift() ingenuo daría rezagos falsos, porque el registro anterior de una celda no
siempre es el mes anterior.
""")

# --- 7. Distribuciones ------------------------------------------------------
s = nueva(prs)
titulo(s, "Distribuciones de las variables")
imagen(s, "14_histogramas_variables.png", 0.7, 1.3, 6.4)
vinetas(s, [
    "Todas unimodales y casi simétricas: la asimetría va de −0.14 a 0.09.",
    "Media y mediana casi coinciden en cada variable.",
    "No hay señal de mezcla de poblaciones: ningún segundo máximo.",
    "La humedad del suelo es la más concentrada (IQR 0.95) pero con el rango más "
    "amplio (10.49): casi todo cerca de lo normal y unos pocos valores muy lejos.",
], izquierda=7.4, arriba=1.6, ancho=Inches(5.4), tam=14)
mensaje(s, "Distribuciones limpias: es lo esperable de índices estandarizados.")
notas(s, E2, 40, """
Diapositiva rápida, no hace falta detenerse.

El mensaje es que no hubo sorpresas: las distribuciones se comportan como se espera
de índices estandarizados, sin colas patológicas ni bimodalidad.

La única que vale señalar es la humedad del suelo, abajo a la izquierda: pico
central mucho más alto y estrecho. Eso anticipa la siguiente diapositiva, porque
será la variable con más atípicos.
""")

# --- 8. Outliers ------------------------------------------------------------
s = nueva(prs)
titulo(s, "Valores atípicos: por qué no los eliminamos")
vinetas(s, [
    "Por criterio IQR: 0.25 % en la variable objetivo, 2.58 % en humedad del suelo.",
    "Decisión: no se eliminó ninguno. Los extremos llegan a −3.97 y 4.06, magnitudes "
    "plausibles; no hay valores imposibles ni códigos centinela como −9999.",
    "La razón de fondo se ve abajo: no se reparten al azar. Se agrupan en años y en "
    "regiones concretos, que es la firma de eventos reales, no de errores de medición.",
], arriba=1.25, tam=15)
imagen(s, "15_atipicos_tiempo_espacio.png", 0.9, 3.25, 11.5)
mensaje(s, "Eliminar los atípicos habría sido borrar el fenómeno de interés.")
notas(s, E2, 60, """
Es una diapositiva de criterio, no de resultado. La rúbrica pide explicar los
atípicos, y la respuesta correcta aquí es no tocarlos.

El argumento más fuerte es el segundo: si fueran errores de medición, estarían
repartidos al azar. Se agrupan en años y regiones concretos, y eso es lo que hace
un evento hidrológico real.

Mencionar también la limitación del método IQR: marca lo que está lejos del grueso
de la distribución, no lo que está mal. Por eso la humedad del suelo, que es muy
concentrada, sale con diez veces más atípicos sin que eso signifique más errores.

Aquí termina el segundo bloque: dar paso al siguiente expositor.
""")

# =============================================================================
# BLOQUE 3 — Patrones temporales
# =============================================================================

# --- 9. Tendencia -----------------------------------------------------------
s = nueva(prs)
titulo(s, "Patrones temporales: una tendencia descendente")
vinetas(s, [
    "El promedio anual cae de 0.322 en 2004 a −0.136 en 2015.",
    "La serie no es estacionaria en media: prueba de Dickey-Fuller con p = 0.526, "
    "no se rechaza la raíz unitaria.",
    "No depende de cómo se agregue: el promedio simple y el ponderado por área "
    "correlacionan 0.936.",
], arriba=1.25, tam=15)
imagen(s, "16_media_movil.png", 1.6, 2.75, 10.1)
mensaje(s, "El registro muestra un descenso sostenido, y la serie tiene memoria larga.")
notas(s, E3, 65, """
Señalar la línea roja: es la media móvil de 12 meses. Explicar por qué 12 y no otro
número — los datos son mensuales y el ciclo esperado es anual, así que una ventana
de 12 promedia un ciclo completo y cancela la estacionalidad, dejando ver la
tendencia.

La caída es visible a simple vista: arranca alrededor de 0.3 en 2004 y termina en
negativo.

Explicar el Dickey-Fuller en términos sencillos: es una prueba de si la serie tiene
un nivel estable al que regresa. Con p = 0.53 no podemos afirmar que lo tenga, lo
cual es coherente con ver una tendencia.

Anticipar la consecuencia: para modelar habrá que diferenciar la serie o incorporar
la tendencia de forma explícita.
""")

# --- 10. Descomposición -----------------------------------------------------
s = nueva(prs)
titulo(s, "Descomposición: ¿dónde está la variabilidad?")
imagen(s, "17_descomposicion.png", 0.7, 1.3, 7.0)
vinetas(s, [
    "La serie se separa en tendencia, estacionalidad y residuo.",
    "Tendencia:  64 % de la varianza",
    "Residuo:  18 %",
    "Estacionalidad:  3 %",
    "El ciclo anual casi no aparece en la serie global. Ese resultado nos "
    "sorprendió y motivó la siguiente diapositiva.",
], izquierda=7.9, arriba=1.7, ancho=Inches(4.9), tam=15)
mensaje(s, "La tendencia manda; el ciclo anual, en apariencia, casi no existe.")
notas(s, E3, 60, """
Explicar primero qué es una descomposición, en una frase por componente:
la tendencia es el movimiento de fondo, la estacionalidad el patrón que se repite
cada 12 meses, y el residuo lo que no explican las otras dos.

Luego dar los porcentajes. El dato llamativo es el 3 % de estacionalidad: en una
variable hidrológica uno esperaría un ciclo anual fuerte.

No resolver el misterio todavía. Plantearlo como pregunta abierta y pasar a la
siguiente diapositiva, que es donde se explica.

Nota técnica por si preguntan: para descomponer hubo que rellenar los 22 meses
ausentes por interpolación, porque la función no admite huecos. Esa interpolación
se usó solo aquí y no se propagó al resto del análisis.
""")

# --- 11. Estacionalidad -----------------------------------------------------
s = nueva(prs)
titulo(s, "La estacionalidad que casi no vemos")
vinetas(s, [
    "La causa: los hemisferios tienen ciclos opuestos y se cancelan al promediarse.",
    "El hemisferio norte promedia positivo los doce meses; el sur cae hasta −0.105 "
    "en junio.",
    "Además el conjunto está desbalanceado: el 80.3 % de las celdas están en el norte.",
], arriba=1.25, tam=15)
imagen(s, "17_perfil_estacional.png", 1.4, 2.9, 10.5)
mensaje(s, "Promediar el planeta entero borra el fenómeno. Hay que mirar por región.")
notas(s, E3, 80, """
Esta es la diapositiva más importante del bloque y la lección metodológica del
proyecto. Darle tiempo.

Señalar las dos curvas de la izquierda: la azul es el norte, la naranja el sur. Van
en direcciones opuestas. Cuando se promedian, se cancelan, y por eso la
descomposición global mostraba apenas 3 % de estacionalidad.

La conclusión que debe quedar: no es que el agua no tenga ciclo anual. Es que una
serie agregada globalmente no es el objeto adecuado para estudiarlo.

Agregar el matiz del desbalance: como el 80 % de las celdas están en el norte, la
serie "global" es en buena medida la serie del hemisferio norte.

Si hay tiempo, mencionar que al cruzar banda de latitud contra mes el ciclo anual
reaparece con claridad dentro de cada región.
""")

# --- 12. Regiones -----------------------------------------------------------
s = nueva(prs)
titulo(s, "Cada región se comporta distinto")
imagen(s, "18_series_por_banda.png", 0.7, 1.35, 7.2)
vinetas(s, [
    "Tropical norte: media −0.222, la más baja.",
    "Subtropical norte: 0.327, la más alta.",
    "Entre ambas hay media desviación estándar de diferencia.",
    "La banda tropical norte concentra además el mínimo (−3.97) y el máximo (4.06) "
    "absolutos: es donde ocurren las anomalías más pronunciadas en ambos sentidos.",
    "La banda polar sur no tiene ninguna observación: la Antártida queda fuera.",
], izquierda=8.1, arriba=1.7, ancho=Inches(4.7), tam=13)
mensaje(s, "El promedio global no representa a ninguna región en particular.")
notas(s, E3, 50, """
Refuerza el mensaje anterior con evidencia espacial.

La idea: no solo los hemisferios difieren, también las bandas de latitud dentro de
un mismo hemisferio. Media desviación estándar de diferencia entre la banda
tropical norte y la subtropical norte es mucho.

El detalle interesante es que la banda tropical norte es a la vez la de menor media
y la de mayor rango: la región más seca en promedio es también donde ocurren los
eventos más extremos.

Aquí termina el tercer bloque: dar paso al último expositor.
""")

# =============================================================================
# BLOQUE 4 — Relaciones y cierre
# =============================================================================

# --- 13. Correlaciones ------------------------------------------------------
s = nueva(prs)
titulo(s, "Principales relaciones entre variables")
imagen(s, "19_matriz_correlacion.png", 0.7, 1.3, 5.4)
vinetas(s, [
    "Correlación con la variable objetivo:",
    "    TWS del mes actual  ·  0.803",
    "    SPEI a 12 meses  ·  0.380",
    "    SPEI a 6 meses  ·  0.368",
    "    Humedad del suelo  ·  0.320",
    "    SPEI a 3 meses  ·  0.285",
    "    SPEI a 1 mes  ·  0.204",
    "Los SPEI se ordenan exactamente por escala de acumulación.",
    "Relaciones lineales: Pearson y Spearman difieren 0.024 como máximo.",
], izquierda=6.4, arriba=1.45, ancho=Inches(6.3), tam=14)
mensaje(s, "Asociación, no causalidad: que se muevan juntas no prueba que una determine a la otra.")
notas(s, E4, 65, """
Dos mensajes.

El primero: la persistencia domina. El mejor predictor del agua del mes siguiente es
el agua del mes actual, con 0.803, más del doble que cualquier covariable. El agua
almacenada tiene mucha inercia.

El segundo es el más interesante y conviene contarlo bien: los cuatro SPEI se
ordenan exactamente por su escala de acumulación, de 12 meses a 1 mes. Nosotros
habíamos planteado esa expectativa antes de mirar los datos, razonando que el TWS
integra agua subterránea, que responde lento. Los datos la confirmaron. Fue una
hipótesis verificada, no una suposición.

Cerrar con la advertencia de causalidad, que es explícita en la rúbrica.
""")

# --- 14. Línea base ---------------------------------------------------------
s = nueva(prs)
titulo(s, "La línea base que hay que superar")
vinetas(s, [
    "Si el agua tiene tanta inercia, la regla más simple ya predice bien: "
    "«el mes que viene será como este».",
    "Esa regla se llama línea base de persistencia y da RMSE 0.5724, equivalente al "
    "62.8 % de la desviación estándar de la objetivo.",
    "Como la métrica del reto es el RMSE, ese número es la referencia concreta: un "
    "modelo que no baje de 0.5724 no aporta nada.",
], arriba=1.25, tam=15)
imagen(s, "19_linea_base_persistencia.png", 1.4, 2.9, 10.5)
mensaje(s, "Todo modelo futuro se mide contra este número.")
notas(s, E4, 55, """
Es la diapositiva que conecta el EDA con la etapa de modelado.

Explicar la idea de línea base: antes de construir algo complejo hay que saber qué
tan lejos llega lo trivial. Aquí lo trivial llega bastante lejos, porque el sistema
tiene inercia.

Señalar la gráfica de la derecha: el RMSE de la persistencia mes a mes. No es
constante, lo que sugiere que hay periodos más difíciles de predecir que otros.

El número a recordar es 0.5724. Si el modelo del próximo proyecto no baja de ahí,
no sirve.
""")

# --- 15. Hallazgos ----------------------------------------------------------
s = nueva(prs)
titulo(s, "Hallazgos más importantes")
vinetas(s, [
    "Sin nulos ni duplicados, pero faltan 22 meses completos del calendario.",
    "La variable objetivo es el TWS del mes calendario t+1 — lo verificamos celda "
    "por celda, no lo dimos por supuesto.",
    "Las variables están estandarizadas: no admiten lectura como volumen físico.",
    "Tendencia descendente y serie no estacionaria en media.",
    "La estacionalidad se cancela globalmente, pero existe por región.",
    "Desbalance espacial: 80 % de las observaciones en el hemisferio norte.",
    "La persistencia domina, con correlación 0.803.",
    "Línea base a superar: RMSE 0.5724.",
], tam=15)
notas(s, E4, 50, """
No leer los ocho puntos: la audiencia los tiene en pantalla. Elegir tres y
comentarlos.

Los que más valor tienen: el de los meses faltantes, porque es el que no detecta
ninguna herramienta automática; el de la verificación de la variable objetivo,
porque muestra que no confiamos en la documentación sino que lo comprobamos; y el
de la estacionalidad, por la lección metodológica.

Mencionar que cada hallazgo está respaldado por una tabla o una figura del
notebook, que es público en el repositorio.
""")

# --- 16. Conclusiones y próximos pasos --------------------------------------
s = nueva(prs)
titulo(s, "Conclusiones y próximos pasos")
vinetas(s, [
    "El problema de calidad no eran los nulos, era la cobertura del calendario.",
    "El promedio global no representa a ninguna región: hay que analizar por zona.",
    "La inercia del sistema es la señal más fuerte disponible.",
    "El análisis es descriptivo y asociativo: no establecemos causalidad.",
], arriba=1.3, ancho=Inches(5.8), tam=15)
vinetas(s, [
    "Reportar siempre contra la línea base (RMSE 0.5724).",
    "Construir rezagos respetando los meses ausentes.",
    "Incorporar la dimensión regional y validar por regiones.",
    "Tratar la tendencia de forma explícita.",
    "Priorizar los SPEI de escala larga (12 y 6 meses).",
    "Validación cronológica, nunca aleatoria.",
    "Manejar el enmascaramiento del 66.53 % de TWS_t en el conjunto de prueba.",
], izquierda=7.1, arriba=1.3, ancho=Inches(5.7), tam=14)
for etiqueta, izq in [("Conclusiones", 0.7), ("Próximos pasos", 7.1)]:
    caja = s.shapes.add_textbox(Inches(izq), Inches(1.0), Inches(4), Inches(0.35))
    p = caja.text_frame.paragraphs[0]
    p.text = etiqueta
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = TEAL
mensaje(s, "Gracias. ¿Preguntas?")
notas(s, E4, 60, """
Cerrar conectando con los objetivos del inicio: los seis se cumplieron y cada uno
se retoma explícitamente en las conclusiones del informe.

De los próximos pasos, destacar dos. El del enmascaramiento es el más delicado: en
el conjunto de prueba, dos tercios de las filas no tienen el TWS del mes actual, y
esa es justamente la variable más predictiva. El modelo tendrá que funcionar sin
ella la mayor parte del tiempo, y hay que reproducir esa condición al validar.

El otro es el de los rezagos, que enlaza con el hallazgo de los meses faltantes.

Terminar agradeciendo y abriendo a preguntas.
""")

# =============================================================================
# Guardado y resumen de tiempos
# =============================================================================
DESTINO.parent.mkdir(parents=True, exist_ok=True)
prs.save(DESTINO)

total = sum(s for _, _, s in guion)
print(f"Presentación creada: {DESTINO.relative_to(RAIZ)}  ({len(guion)} diapositivas)\n")
print(f"  {'#':>2}  {'Expositor':<14} {'Tiempo':>7}")
print("  " + "-" * 28)
for numero, expositor, segundos in guion:
    print(f"  {numero:>2}  {expositor:<14} {segundos // 60}:{segundos % 60:02d}".rstrip())
print("  " + "-" * 28)
print(f"  {'TOTAL':<18} {total // 60}:{total % 60:02d}\n")

for expositor in (E1, E2, E3, E4):
    suyo = sum(s for _, e, s in guion if e == expositor)
    laminas = [str(n) for n, e, _ in guion if e == expositor]
    print(f"  {expositor:<14} {suyo // 60}:{suyo % 60:02d}   (diapositivas {', '.join(laminas)})")
