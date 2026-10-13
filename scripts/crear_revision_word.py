"""Crea la revisión bibliográfica editable en Word, con formato negro."""

from __future__ import annotations

import re
from pathlib import Path

from lxml import etree

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs/Revision_Bibliografica_TWS.md"
OUTPUT = ROOT / "output/word/Revision_Bibliografica_y_Seleccion_Modelos_TWS.docx"
BLACK = RGBColor(0, 0, 0)
TOKEN = re.compile(r"\[([^\]]+)\]\((https?://[^\s)]+)\)|\*\*(.+?)\*\*|\*([^*]+)\*|`([^`]+)`")


def black_run(run, *, size=None):
    run.font.name = "Arial"
    run.font.color.rgb = BLACK
    if size is not None:
        run.font.size = Pt(size)
    return run


def hyperlink(paragraph, label, url):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    props.append(style)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def rich(paragraph, text, *, size=None):
    pos = 0
    for match in TOKEN.finditer(text):
        if match.start() > pos:
            black_run(paragraph.add_run(text[pos:match.start()]), size=size)
        if match[1] is not None:
            hyperlink(paragraph, match[1], match[2])
        else:
            content = match[3] or match[4] or match[5]
            run = black_run(paragraph.add_run(content), size=size)
            run.bold = match[3] is not None
            run.italic = match[4] is not None
        pos = match.end()
    if pos < len(text):
        black_run(paragraph.add_run(text[pos:]), size=size)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.95)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)
    section.header_distance = Cm(0.85)
    section.footer_distance = Cm(0.85)
    for style in doc.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            style.font.name = "Arial"
            style.font.color.rgb = BLACK
            for color in style.element.xpath(".//w:color"):
                color.attrib.clear()
                color.set(qn("w:val"), "000000")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.8)
    normal.paragraph_format.line_spacing = Pt(15)
    normal.paragraph_format.space_after = Pt(9)
    normal.paragraph_format.widow_control = True
    for name, size, spacing in [("Title", 25, 15), ("Subtitle", 16, 22),
                                ("Heading 1", 17, 17), ("Heading 2", 12.5, 9),
                                ("Heading 3", 11.2, 6)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(5 if name == "Heading 3" else 0)
        style.paragraph_format.space_after = Pt(spacing)
        style.paragraph_format.line_spacing = Pt(size * 1.24)
        style.paragraph_format.keep_with_next = True
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    for name in ("Hyperlink", "FollowedHyperlink"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
        doc.styles[name].font.color.rgb = BLACK
        doc.styles[name].font.underline = False
    for name, size, leading in [("Reference", 10.6, 20.8), ("Note", 9.2, 12.7),
                                ("Table Text", 8.9, 12), ("Metadata", 10.8, 15)]:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.paragraph_format.line_spacing = Pt(leading)
        style.paragraph_format.space_after = Pt(10 if name == "Reference" else 8)
    ref = doc.styles["Reference"].paragraph_format
    ref.left_indent = Cm(1.27)
    ref.first_line_indent = Cm(-1.27)
    ref.keep_together = True
    table = doc.styles["Table Text"].paragraph_format
    table.space_after = Pt(0)
    table.keep_together = True
    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    black_run(header.add_run("CC3084 · PROYECTO 2"), size=8)
    header.paragraph_format.tab_stops.add_tab_stop(Cm(16.7))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    black_run(footer.add_run("Página "), size=8)
    field_run = black_run(footer.add_run(), size=8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    cached = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    cached.append(t)
    field.append(cached)
    field_run._r.addnext(field)
    doc.core_properties.title = "Revisión bibliográfica y selección preliminar de modelos de TWS"
    doc.core_properties.subject = "Pronóstico del almacenamiento total de agua con datos de GRACE"
    doc.core_properties.author = "Proyecto 2 | CC3084 - Data Science"
    doc.core_properties.keywords = "GRACE, TWS, revisión bibliográfica, aprendizaje automático"


def add_table(doc, lines):
    parsed = [[x.strip() for x in line.strip("|").split("|")] for line in lines]
    rows = [parsed[0]] + parsed[2:]
    first = rows[0][0] == "Estudio"
    fractions = [0.205, 0.235, 0.255, 0.305] if first else [0.215, 0.275, 0.20, 0.31]
    table = doc.add_table(rows=len(rows), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    width = 16.7
    for column, fraction in zip(table.columns, fractions):
        column.width = Cm(width * fraction)
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        for key, value in {"val": "single", "sz": "4", "color": "000000"}.items():
            el.set(qn("w:" + key), value)
        borders.append(el)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", 100), ("bottom", 100), ("left", 110), ("right", 110)):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    props.append(margins)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, values in enumerate(rows):
        row = table.rows[i]
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Cm(width * fractions[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = "Table Text"
            rich(p, value)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(5)
    spacer.paragraph_format.space_before = Pt(0)
    black_run(spacer.add_run(), size=1)


def build():
    doc = Document()
    configure(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    page = 1
    references = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "<!-- pagebreak -->":
            page += 1
            # Place the break on the next paragraph, rather than inserting
            # a spare paragraph that could produce an empty page.
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1
            line = lines[i].strip()
            page_break = True
        else:
            page_break = False
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, table_lines)
            continue
        match = re.match(r"^(#{1,3})\s+(.*)$", line)
        if match:
            level, text = len(match[1]), match[2]
            if text.startswith("8. Referencias") or text.startswith("Referencias (continuación)"):
                references = True
            elif text.startswith("Anexo."):
                references = False
            style = "Title" if page == 1 and level == 1 else "Subtitle" if page == 1 and level == 2 else f"Heading {level}"
            p = doc.add_paragraph(style=style)
            rich(p, text)
        else:
            # The source uses fixed PDF page references; retain navigation
            # without stale pagination when Word reflows the text.
            if line.startswith("Problema y búsqueda bibliográfica (p."):
                line = "Problema y búsqueda bibliográfica; antecedentes y síntesis crítica; selección y evaluación propuestas; referencias APA 7; verificación de indexación."
            is_note = line.startswith("*Nota")
            is_caption = line.startswith("**Tabla")
            style = "Note" if is_note else "Reference" if references else "Metadata" if page == 1 and i < 12 else "Normal"
            p = doc.add_paragraph(style=style)
            if style == "Normal":
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if is_caption:
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(8)
            rich(p, line)
        if page_break:
            p.paragraph_format.page_break_before = True
        i += 1
    # Remove all theme/direct color overrides in document and styles,
    # including hyperlinks and headers, so Word cannot restore blue text.
    for part in doc.part.package.parts:
        element = getattr(part, "element", None)
        if element is None:
            # Word includes a second style definition part that python-docx
            # keeps as raw XML. Normalize that fallback as well.
            if str(part.partname) == "/word/stylesWithEffects.xml":
                root = etree.fromstring(part.blob)
                for color in root.iter(qn("w:color")):
                    color.attrib.clear()
                    color.set(qn("w:val"), "000000")
                part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            continue
        for color in element.xpath(".//w:color"):
            color.attrib.clear()
            color.set(qn("w:val"), "000000")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
